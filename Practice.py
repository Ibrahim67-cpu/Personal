from docx import Document

# Create a blank document
doc = Document()

# Add text
doc.add_heading('Success!', level=1)
doc.add_paragraph('python-docx is working perfectly on my Mac.')

# Save it
doc.save('test.docx')
print("File 'test.docx' created successfully!")

